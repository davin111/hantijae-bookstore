"""한티재 ERP SA408/409 XLSX 검산. 외부 연결·DB·LLM 없이 원본의 두 줄 헤더를 읽는다."""
import calendar
import re
import unicodedata
from datetime import date
from io import BytesIO
from zipfile import ZipFile

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10000
METRICS = ('out_qty', 'out_amount', 'return_qty', 'return_amount', 'net_qty', 'net_amount', 'gift_qty')
CUSTOMER_TOP = ['코드', '지역', '거래처명', '출고', '', '출고금액', '반품', '', '출고금액', '순매출', '', '출고금액', '기증수량']
CUSTOMER_SUB = ['', '', '', '수량', '금액', '건수', '수량', '금액', '건수', '수량', '금액', '건수', '수량']
PRODUCT_TOP = ['코드', '제품명', '정가', '출고', '출고금액', '반품', '반품금액', '순출고', '판매금액', '기증수량']
PRODUCT_SUB = ['', '', '', '수량', '금액', '수량', '금액', '수량', '금액', '']


class InvalidERP(ValueError):
    """운영진에게 그대로 보여도 되는 짧은 검산 실패 이유."""


def norm(value):
    return unicodedata.normalize('NFC', str(value or '')).strip()


def _number(value):
    if value is None or value == '':
        return 0
    s = str(value).strip()
    if not re.fullmatch(r'-?(?:\d+|\d{1,3}(?:,\d{3})+)(?:\.0+)?', s):
        raise InvalidERP('수량·금액·건수에 정수가 아닌 값이 있어요.')
    n = int(s.replace(',', '').split('.')[0])
    if abs(n) > 10**12:
        raise InvalidERP('수량·금액의 범위를 확인해 주세요.')
    return n


def _values(row, kind):
    columns = (3, 4, 6, 7, 9, 10, 12) if kind == 'customer' else (3, 4, 5, 6, 7, 8, 9)
    v = dict(zip(METRICS, (_number(row[i]) for i in columns)))
    if any(v[k] < 0 for k in ('out_qty', 'out_amount', 'return_qty', 'return_amount', 'gift_qty')):
        raise InvalidERP('출고·반품·기증 칸의 음수를 확인해 주세요.')
    if v['out_qty'] - v['return_qty'] != v['net_qty'] or v['out_amount'] - v['return_amount'] != v['net_amount']:
        raise InvalidERP('출고에서 반품을 뺀 값이 순출고와 달라요.')
    if kind == 'customer':
        out_count, return_count, net_count = (_number(row[i]) for i in (5, 8, 11))
        if min(out_count, return_count) < 0 or out_count - return_count != net_count:
            raise InvalidERP('거래처별 건수의 차감 계산이 맞지 않아요.')
    return v


def parse_export(data, today):
    if len(data) > MAX_BYTES:
        raise InvalidERP('파일이 너무 커요. 월별 XLSX 파일(5MB 이하)로 보내 주세요.')
    try:
        with ZipFile(BytesIO(data)) as z:
            parts = z.infolist()
            if len(parts) > 300 or sum(p.file_size for p in parts) > 30 * 1024 * 1024:
                raise InvalidERP('압축을 푼 엑셀 파일이 너무 커요.')
            if any(p.flag_bits & 1 or 'vbaProject' in p.filename for p in parts):
                raise InvalidERP('암호·매크로 없는 XLSX 파일로 보내 주세요.')
        # Imports stay here: the status page must not load spreadsheet libraries.
        from openpyxl import load_workbook
        wb = load_workbook(BytesIO(data), read_only=True, data_only=False, keep_links=False)
    except InvalidERP:
        raise
    except Exception:
        raise InvalidERP('XLSX 파일을 읽지 못했어요. ERP에서 다시 내보내 주세요.') from None
    try:
        if len(wb.worksheets) != 1:
            raise InvalidERP('집계표 하나만 있는 XLSX 파일로 보내 주세요.')
        ws = wb.worksheets[0]
        if (ws.max_column or 0) > 13 or (ws.max_row or 0) > MAX_ROWS + 8:
            raise InvalidERP('집계표의 행·열 범위가 예상과 달라요.')
        ws.reset_dimensions()
        rows = []
        for row in ws.iter_rows(max_row=MAX_ROWS + 9, max_col=13, values_only=True):
            if any(isinstance(v, str) and v.startswith('=') for v in row):
                raise InvalidERP('수식이 들어 있어요. ERP에서 값으로 내보낸 원본을 보내 주세요.')
            rows.append(list(row))
        while rows and not any(v is not None and v != '' for v in rows[-1]):
            rows.pop()
        if len(rows) < 8 or len(rows) > MAX_ROWS + 8:
            raise InvalidERP('집계표의 행 수가 예상과 달라요.')
        title = re.sub(r'\s+', '', norm(rows[0][0]))
        kind = {'거래처집계현황-SA409': 'customer', '제품별집계현황-SA408': 'product'}.get(title)
        if not kind:
            raise InvalidERP('거래처 집계 현황(SA409)·제품별 집계 현황(SA408) 파일이 필요해요.')
        if re.sub(r'\s+', '', norm(rows[1][0])) != '회사명:도서출판한티재':
            raise InvalidERP('회사명이 도서출판 한티재인 집계표가 필요해요.')
        header = ' '.join(norm(v) for row in rows[:5] for v in row if v is not None)
        product_filters = [norm(v) for row in rows[:5] for v in row if '제품코드' in norm(v)]
        if len(product_filters) != 1 or not re.match(r'^제품코드\s*:\s*-\s*(?:$|부\s*가\s*세\s*:)', product_filters[0]):
            raise InvalidERP('제품코드 범위가 제한됐거나 중복 표기돼 있어요. 전체 범위로 내보내 주세요.')
        periods = re.findall(r'조회기간\s*:\s*(\d{4}-\d{2}-\d{2})\s*~\s*(\d{4}-\d{2}-\d{2})', header)
        if len(periods) != 1 or '조회기간' not in norm(rows[4][0]):
            raise InvalidERP('조회기간이 하나로 확인되지 않아요.')
        try:
            start, end = (date.fromisoformat(x) for x in periods[0])
        except (ValueError, AttributeError):
            raise InvalidERP('조회기간을 확인하지 못했어요.') from None
        if start.day != 1 or end != date(start.year, start.month, calendar.monthrange(start.year, start.month)[1]) or end >= today:
            raise InvalidERP('마감한 달의 1일부터 말일까지 조회한 파일이 필요해요.')
        if not re.search(r'부\s*가\s*세\s*:\s*별도', header):
            raise InvalidERP('부가세 별도 집계인지 확인해 주세요.')
        top, sub = (CUSTOMER_TOP, CUSTOMER_SUB) if kind == 'customer' else (PRODUCT_TOP, PRODUCT_SUB)
        if [norm(x) for x in rows[5]] != top + [''] * (13 - len(top)) or [norm(x) for x in rows[6]] != sub + [''] * (13 - len(sub)):
            raise InvalidERP('집계표의 열 형식이 달라요. 두 줄 제목을 유지한 원본으로 보내 주세요.')
        details, footer, codes, counts = [], None, set(), [0, 0, 0]
        for row in rows[7:]:
            if not any(v is not None and v != '' for v in row):
                continue
            if footer is not None:
                raise InvalidERP('합계 행 뒤에 다른 내용이 있어요.')
            if not any(norm(v) for v in row[:3]):
                footer = _values(row, kind)
                if kind == 'customer' and [_number(row[i]) for i in (5, 8, 11)] != counts:
                    raise InvalidERP('거래처별 건수와 합계 행이 달라요.')
                continue
            code, label = norm(row[0]), norm(row[2] if kind == 'customer' else row[1])
            if not re.fullmatch(r'\d{6}' if kind == 'customer' else r'\d{8}', code) or not label or code in codes:
                raise InvalidERP('코드·이름이 비었거나 같은 코드가 중복돼 있어요.')
            codes.add(code)
            if kind == 'product':
                if _number(row[2]) < 0:
                    raise InvalidERP('제품의 정가가 음수예요.')
            else:
                counts = [a + _number(row[i]) for a, i in zip(counts, (5, 8, 11))]
            details.append(dict(code=code, label=label, **_values(row, kind)))
        totals = {k: sum(row[k] for row in details) for k in METRICS}
        if footer is None or footer != totals:
            raise InvalidERP('상세 행을 더한 값과 파일의 합계가 달라요.')
        return dict(kind=kind, month=start.strftime('%Y-%m'), totals=totals, rows=details,
                    count=len(details), vat='별도', scope='조회 조건 미확인')
    except InvalidERP:
        raise
    except Exception:
        raise InvalidERP('집계표의 내용을 읽지 못했어요. 원본 형식을 확인해 주세요.') from None
    finally:
        wb.close()


def validate_pair(customer, product):
    if customer['kind'] != 'customer' or product['kind'] != 'product':
        raise InvalidERP('거래처별·제품별 두 종류의 파일이 필요해요.')
    if customer['month'] != product['month'] or customer['vat'] != product['vat']:
        raise InvalidERP('두 파일의 조회기간·부가세 기준이 달라요.')
    if customer['totals'] != product['totals']:
        raise InvalidERP('거래처별·제품별 합계가 달라요. 같은 조회 조건으로 두 파일을 다시 보내 주세요.')
