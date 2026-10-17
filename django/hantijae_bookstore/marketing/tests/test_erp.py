"""Synthetic ERP exports: no customer/business originals in the public repository."""
from datetime import date
from io import BytesIO
from unittest import TestCase

from openpyxl import Workbook

from marketing.erp_parser import InvalidERP, parse_export, validate_pair


def export(kind='customer', month='2026-09', delta=0, change=None):
    wb = Workbook()
    ws = wb.active
    ws.append(['거래처 집계 현황 - SA409' if kind == 'customer' else '제품별 집계 현황 - SA408'])
    ws.append(['회 사 명 : 도서출판 한티재'])
    ws.append(['제품코드 :  -         부 가 세 : 별도'])
    ws.append(['편집부서 :         손익부서 :         마케팅부서 : '])
    ws.append([f'조회기간 : {month}-01 ~ {month}-30' + (' 집계기준 : 거래처 부 가 세 : 별도' if kind == 'customer' else '')])
    if kind == 'customer':
        ws.append(['코드', '지역', '거래처명', '출고', None, '출고금액', '반품', None, '출고금액', '순매출', None, '출고금액', '기증수량'])
        ws.append([None, None, None, '수량', '금액', '건수', '수량', '금액', '건수', '수량', '금액', '건수', '수량'])
        ws.append(['000001', '서울', '가상서점', 4 + delta, 4000 + delta * 1000, 2, 1, 1000, 1, 3 + delta, 3000 + delta * 1000, 1, 2])
        ws.append(['', '', '', str(4 + delta), f'{4000 + delta * 1000:,}', '2', '1', '1,000', '1', str(3 + delta), f'{3000 + delta * 1000:,}', '1', '2'])
    else:
        ws.append(['코드', '제품명', '정가', '출고', '출고금액', '반품', '반품금액', '순출고', '판매금액', '기증수량'])
        ws.append([None, None, None, '수량', '금액', '수량', '금액', '수량', '금액', None])
        ws.append(['00000001', '가상 책', '2,000', 4 + delta, 4000 + delta * 1000, 0, 0, 4 + delta, 4000 + delta * 1000, 2])
        ws.append(['00000002', '반품 책', '2,000', 0, 0, 1, 1000, -1, -1000, 0])
        ws.append(['', '', '', str(4 + delta), f'{4000 + delta * 1000:,}', '1', '1,000', str(3 + delta), f'{3000 + delta * 1000:,}', '2'])
    if change:
        change(ws)
    out = BytesIO()
    wb.save(out)
    wb.close()
    return out.getvalue()


class ERPParserTest(TestCase):
    def read(self, kind='customer', **kwargs):
        return parse_export(export(kind, **kwargs), today=date(2026, 10, 17))

    def test_two_level_header_count_columns_are_not_money(self):
        customer, product = self.read(), self.read('product')
        validate_pair(customer, product)
        self.assertEqual(customer['totals'], dict(out_qty=4, out_amount=4000, return_qty=1,
                         return_amount=1000, net_qty=3, net_amount=3000, gift_qty=2))
        self.assertEqual(product['rows'][1]['net_qty'], -1)
        self.assertEqual(customer['rows'][0]['code'], '000001')
        self.assertEqual((customer['count'], product['count']), (1, 2))
        self.assertEqual(customer['scope'], '조회 조건 미확인')

    def test_changed_column_is_rejected(self):
        with self.assertRaisesRegex(InvalidERP, '열'):
            self.read(change=lambda ws: setattr(ws['E7'], 'value', '건수'))

    def test_row_and_footer_must_both_reconcile(self):
        for cell, value in [('J8', 5), ('K9', '4,000'), ('F8', 9)]:
            with self.subTest(cell=cell), self.assertRaises(InvalidERP):
                self.read(change=lambda ws: setattr(ws[cell], 'value', value))

    def test_formulas_unknown_numbers_and_duplicate_codes_rejected(self):
        for cell, value in [('E8', '=2000*2'), ('E8', '4천'), ('D8', 4.5), ('A10', '000001')]:
            with self.subTest(cell=cell), self.assertRaises(InvalidERP):
                self.read(change=lambda ws: setattr(ws[cell], 'value', value))

    def test_full_closed_calendar_month_required(self):
        for period in ['2026-09-02 ~ 2026-09-30', '2026-09-01 ~ 2026-09-29', '2026-10-01 ~ 2026-10-31']:
            with self.subTest(period=period), self.assertRaises(InvalidERP):
                self.read(change=lambda ws: setattr(ws['A5'], 'value', f'조회기간 : {period} 부 가 세 : 별도'))

    def test_wrong_company_and_extra_sheet_rejected(self):
        with self.assertRaises(InvalidERP):
            self.read(change=lambda ws: setattr(ws['A2'], 'value', '회 사 명 : 다른 출판사'))
        with self.assertRaises(InvalidERP):
            self.read(change=lambda ws: ws.parent.create_sheet('extra'))

    def test_pair_must_match_month_kind_and_all_seven_totals(self):
        c, p = self.read(), self.read('product')
        for changed in [dict(p, month='2026-08'), dict(p, kind='customer'), self.read('product', delta=1)]:
            with self.assertRaises(InvalidERP):
                validate_pair(c, changed)

    def test_non_xlsx_rejected_without_parser_details(self):
        with self.assertRaises(InvalidERP):
            parse_export(b'not a workbook', today=date(2026, 10, 17))

    def test_conflicting_period_or_restricted_product_range_rejected(self):
        for extra in ['제품코드 : 00000001 - 00000002', '조회기간 : 2026-08-01 ~ 2026-08-31']:
            with self.subTest(extra=extra), self.assertRaises(InvalidERP):
                self.read(change=lambda ws: setattr(ws['A4'], 'value', extra))

    def test_negative_price_and_gross_counts_rejected(self):
        with self.assertRaises(InvalidERP):
            self.read('product', change=lambda ws: setattr(ws['C8'], 'value', -2000))
        def counts(ws):
            for row in (8, 9):
                ws.cell(row, 6, -2)
                ws.cell(row, 9, -1)
                ws.cell(row, 12, -1)
        with self.assertRaises(InvalidERP):
            self.read(change=counts)
