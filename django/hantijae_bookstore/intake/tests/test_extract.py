import io
import os
import tempfile
import zipfile
from unittest import mock

import docx
from django.test import SimpleTestCase

from intake import extract

HP = 'http://www.hancom.co.kr/hwpml/2011/paragraph'
SECTION = f'''<?xml version="1.0" encoding="UTF-8"?>
<hs:sec xmlns:hs="http://www.hancom.co.kr/hwpml/2011/section" xmlns:hp="{HP}">
  <hp:p><hp:run><hp:t>[보도자료]</hp:t></hp:run></hp:p>
  <hp:p><hp:run><hp:tbl><hp:tr><hp:tc><hp:subList>
    <hp:p><hp:run><hp:t>판형 </hp:t><hp:t>130×200</hp:t></hp:run></hp:p>
  </hp:subList></hp:tc></hp:tr></hp:tbl></hp:run></hp:p>
  <hp:p><hp:run><hp:t>ISBN 979-11-92455-87-7</hp:t></hp:run></hp:p>
</hs:sec>'''

CLEAN = ('한국 최초로 커밍아웃한 트랜스젠더 변호사인 저자는 자신의 정체성 형성 과정과 이중 생활의 고뇌, '
         '그리고 로스쿨 진학 후 단계적 커밍아웃을 거쳐 활동가로 거듭나기까지의 여정을 들려준다. 지은이 박한희 ISBN 발행일')
GARBLED = ('연대와  환대 지은이박지호 판형 무선 1 3 0 × 1 8 5 면수 쪽 1 3 4 가격 원 1 3 , 0 0 0 ISBN 9 7 9 - 1 1 '
           '발행일 년 월 일 2 0 2 4 1 0 1 명이나 되었다고 한다 그들은 . 보냈다 그중에는 . , 로 부 터 쉽 게')


class ExtractTest(SimpleTestCase):
    def setUp(self):
        self.tmp = tempfile.mkdtemp()

    def test_hwpx_reads_nested_table_paragraphs_once(self):
        path = os.path.join(self.tmp, 'a.hwpx')
        with zipfile.ZipFile(path, 'w') as z:
            z.writestr('Contents/section0.xml', SECTION)
        text = extract.extract_text(path)
        self.assertEqual(text.splitlines(), ['[보도자료]', '판형 130×200', 'ISBN 979-11-92455-87-7'])

    def test_docx(self):
        path = os.path.join(self.tmp, 'a.docx')
        d = docx.Document()
        d.add_paragraph('■ 책 소개')
        d.add_paragraph('본문입니다')
        d.save(path)
        self.assertEqual(extract.extract_text(path).splitlines(), ['■ 책 소개', '본문입니다'])

    def test_pdf_joins_pages(self):
        pages = [mock.Mock(extract_text=lambda: '첫 쪽'), mock.Mock(extract_text=lambda: '둘째 쪽')]
        with mock.patch('intake.extract.pypdf.PdfReader', return_value=mock.Mock(pages=pages)):
            self.assertEqual(extract.extract_text('/x/a.pdf'), '첫 쪽\n둘째 쪽')

    def test_hwp_uses_hwp5txt_and_drops_placeholders(self):
        done = mock.Mock(returncode=0, stdout='[보도자료]\n<그림>\n<표>\n본문\n'.encode())
        with mock.patch('intake.extract.subprocess.run', return_value=done):
            self.assertEqual(extract.extract_text('/x/a.hwp'), '[보도자료]\n본문')

    def test_quality_score_separates_clean_and_garbled(self):
        self.assertGreater(extract.quality_score(CLEAN * 3), 0.8)
        self.assertLess(extract.quality_score(GARBLED * 3), 0.6)
        self.assertEqual(extract.quality_score('짧음'), 0.0)

    def test_best_local_text_falls_through_bad_sources(self):
        def fake(path):
            if path.endswith('.hwpx'):
                raise extract.ExtractError('broken')
            return GARBLED * 3 if path.endswith('.pdf') else CLEAN * 3
        with mock.patch('intake.extract.extract_text', side_effect=fake):
            text, score, used = extract.best_local_text(['a.hwpx', 'a.hwp', 'a.pdf'])
        self.assertEqual(used, 'a.hwp')

    def test_is_verbatim_ignores_whitespace_only(self):
        source = '한국 최초로 커밍아웃한\n트랜스젠더 변호사인 저자는 자신의 정체성 형성 과\n정과'
        self.assertTrue(extract.is_verbatim('한국 최초로 커밍아웃한 트랜스젠더 변호사인 저자는', source))
        self.assertTrue(extract.is_verbatim('자신의 정체성 형성 과정과', source))
        self.assertFalse(extract.is_verbatim('한국 최초로 커밍아웃한 트랜스젠더 변호사 저자는', source))

    def test_realistic_partially_garbled_pdf_is_below_threshold(self):
        # 2024년 PDF: 헤더 숫자만 한 글자씩 떨어지고 본문은 대부분 멀쩡한 형태
        mixed = ('보도자료[] 한티재 팸플릿 028 연대와 환대 지은이박지호 | 분야사회 | 판형 무선 1 3 0 × 1 8 5 | 면수 쪽 1 3 4도(1 ) '
                 '가격 원 1 3 , 0 0 0 | ISBN 9 7 9 - 1 1 - 9 2 4 5 5 - 5 9 - 4 0 4 3 0 0 | 발행일 년 월 일 2 0 2 4 1 0 1 '
                 '년 월 일 밀양행정대집행 년을 맞아 희망버스를 타고 현장에 함께한 사람들이 여 2024 6 8 , 10 1,500 '
                 '명이나 되었다고 한다 그들은 밀양 할매들과 울다가 웃다가 춤추다가 하면서 다섯 시간을 함께 . ')
        realistic = mixed + ' ' + CLEAN * 2   # 실제 파일은 0.63: 헤더는 깨지고 본문은 대체로 멀쩡
        self.assertLess(extract.quality_score(realistic), extract.QUALITY_THRESHOLD)
        self.assertGreater(extract.quality_score(CLEAN * 3), extract.QUALITY_THRESHOLD)
