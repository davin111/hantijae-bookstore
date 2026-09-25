"""보도자료 문서 → 텍스트. 사이드카 입력이 아니라 원문 대조·예비용."""
import os
import re
import shutil
import subprocess
import sys
import unicodedata
import xml.etree.ElementTree as ET
import zipfile
from typing import Sequence, Tuple

import docx
import pypdf

HP = '{http://www.hancom.co.kr/hwpml/2011/paragraph}'
HWP_PLACEHOLDERS = {'<그림>', '<표>', '<수식>'}
# 실측: 깨끗한 2026 PDF·HWPX·HWP 0.84~0.92, 글자가 흩어진 2024 PDF 0.63
QUALITY_THRESHOLD = 0.75


class ExtractError(Exception):
    pass


def _pdf_text(path):
    reader = pypdf.PdfReader(path)
    return '\n'.join((page.extract_text() or '') for page in reader.pages)


def _hwpx_text(path):
    lines = []
    with zipfile.ZipFile(path) as z:
        sections = sorted((n for n in z.namelist() if re.fullmatch(r'Contents/section\d+\.xml', n)),
                          key=lambda n: int(re.search(r'\d+', n).group()))
        for name in sections:
            root = ET.fromstring(z.read(name))
            # iter()는 표 안의 중첩 문단도 순서대로 돈다. 각 문단은 "직속 run의 t"만 모아 중복을 피한다.
            for p in root.iter(HP + 'p'):
                text = ''.join(''.join(t.itertext()) for run in p.findall(HP + 'run') for t in run.findall(HP + 't'))
                if text.strip():
                    lines.append(text)
    return '\n'.join(lines)


def _hwp_text(path):
    exe = os.path.join(os.path.dirname(sys.executable), 'hwp5txt')
    if not os.path.exists(exe):
        exe = shutil.which('hwp5txt') or 'hwp5txt'
    result = subprocess.run([exe, path], capture_output=True, timeout=120)
    if result.returncode != 0:
        raise ExtractError(f'hwp5txt 실패 ({result.returncode})')
    lines = result.stdout.decode('utf-8', errors='replace').splitlines()
    return '\n'.join(l for l in lines if l.strip() and l.strip() not in HWP_PLACEHOLDERS)


def _docx_text(path):
    d = docx.Document(path)
    lines = [p.text for p in d.paragraphs if p.text.strip()]
    for table in d.tables:
        for row in table.rows:
            lines.append(' | '.join(cell.text.strip() for cell in row.cells))
    return '\n'.join(lines)


EXTRACTORS = {'.pdf': _pdf_text, '.hwpx': _hwpx_text, '.hwp': _hwp_text, '.docx': _docx_text}


def extract_text(path: str) -> str:
    ext = os.path.splitext(path)[1].lower()
    if ext not in EXTRACTORS:
        raise ExtractError(f'지원하지 않는 형식: {ext}')
    try:
        return unicodedata.normalize('NFC', EXTRACTORS[ext](path))
    except ExtractError:
        raise
    except Exception as e:  # 손상 파일 등
        raise ExtractError(f'{os.path.basename(path)}: {e}') from e


def quality_score(text: str) -> float:
    """글자가 한 자씩 떨어진 PDF(`1 3 0 × 1 8 5`)를 걸러낸다."""
    tokens = text.split()
    if len(tokens) < 20:
        return 0.0
    singles = sum(1 for t in tokens if len(t) == 1 and ('가' <= t <= '힣' or t.isdigit()))
    ratio = singles / len(tokens)
    markers = sum(k in text for k in ('ISBN', '지은이', '발행일')) / 3
    return round(max(0.0, 1 - ratio * 3) * 0.7 + markers * 0.3, 3)


def best_local_text(paths: Sequence[str], threshold: float = QUALITY_THRESHOLD) -> Tuple[str, float, str]:
    best = ('', 0.0, '')
    for path in paths:
        try:
            text = extract_text(path)
        except ExtractError:
            continue
        score = quality_score(text)
        if score >= threshold:
            return text, score, path
        if score > best[1]:
            best = (text, score, path)
    return best


def compact(s: str) -> str:
    return re.sub(r'\s+', '', unicodedata.normalize('NFC', s or ''))


def is_verbatim(candidate: str, source: str) -> bool:
    c = compact(candidate)
    return bool(c) and c in compact(source)
